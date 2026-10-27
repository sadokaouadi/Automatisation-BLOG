from pathlib import Path
from datetime import datetime, date

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.database.database import get_connection


PROJECT_ROOT = Path(__file__).resolve().parents[2]

BLOG_FILE = (
    PROJECT_ROOT
    / "data"
    / "imports"
    / "BLOG_2026_RJ45.xlsx"
)


def clean_text(value):
    """Transforme PO, UPC, etc. en texte propre."""

    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def clean_date(value, epoch):
    """Transforme les dates Excel en YYYY-MM-DD."""

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, (int, float)):
        try:
            return from_excel(value, epoch).date().isoformat()
        except Exception:
            return str(value)

    return str(value).strip()


def import_rj45_blog():

    if not BLOG_FILE.exists():
        print(f"ERREUR : fichier introuvable : {BLOG_FILE}")
        return

    print("Lecture du BLOG RJ45...")

    workbook = load_workbook(
        BLOG_FILE,
        read_only=True,
        data_only=True
    )

    if "RJ45" not in workbook.sheetnames:
        print("ERREUR : feuille RJ45 introuvable.")
        workbook.close()
        return

    sheet = workbook["RJ45"]

    connection = get_connection()
    cursor = connection.cursor()

    imported = 0
    skipped = 0
    duplicates = 0

    # Ligne 1 = entêtes
    # Colonnes A:G uniquement
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            max_col=7,
            values_only=True
        ),
        start=2
    ):

        (
            order_date,
            customer_date,
            po,
            product,
            quantity,
            length,
            upc
        ) = row

        # Ignore les lignes qui ne sont pas des commandes
        if not po or not product or quantity is None:
            skipped += 1
            continue

        po = clean_text(po)
        product = clean_text(product)
        upc = clean_text(upc)

        order_date = clean_date(
            order_date,
            workbook.epoch
        )

        customer_date = clean_date(
            customer_date,
            workbook.epoch
        )

        try:
            quantity = int(quantity)
        except (ValueError, TypeError):
            skipped += 1
            continue

        try:
            length = float(length) if length is not None else None
        except (ValueError, TypeError):
            length = None

        # Vérifie que la commande n'est pas déjà importée
        cursor.execute(
            """
            SELECT id
            FROM orders
            WHERE po = ?
              AND product = ?
              AND quantity = ?
            LIMIT 1
            """,
            (
                po,
                product,
                quantity
            )
        )

        if cursor.fetchone():
            duplicates += 1
            continue

        cursor.execute(
            """
            INSERT INTO orders (
                date,
                customer_date,
                po,
                product,
                quantity,
                length,
                upc,
                customer,
                status,
                verification_1,
                verification_2
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                order_date,
                customer_date,
                po,
                product,
                quantity,
                length,
                upc,
                None,
                "HISTORIQUE",
                1,
                1
            )
        )

        imported += 1

    connection.commit()

    connection.close()
    workbook.close()

    print()
    print("IMPORT TERMINE")
    print("----------------------------")
    print(f"Commandes importees : {imported}")
    print(f"Lignes ignorees     : {skipped}")
    print(f"Doublons ignores    : {duplicates}")
    print("----------------------------")


if __name__ == "__main__":
    import_rj45_blog()