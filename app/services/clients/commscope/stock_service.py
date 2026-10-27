from pathlib import Path
from functools import lru_cache
from threading import Lock

from openpyxl import load_workbook

from app.services.clients.commscope.rj45_stock_analyzer import (
    find_component_blocks,
    find_last_order_row
)


PROJECT_ROOT = Path(__file__).resolve().parents[4]

BLOG_PATH = (
    PROJECT_ROOT
    / "data"
    / "imports"
    / "BLOG_2026_RJ45.xlsx"
)

_stock_cache_lock = Lock()


@lru_cache(maxsize=2)
def _load_current_stock(path, modified_time, file_size):
    """
    Retourne le stock actuel de tous les composants
    du BLOG RJ45 CommScope.
    """

    blog_path = Path(path)
    if not blog_path.exists():
        raise FileNotFoundError(
            f"BLOG RJ45 introuvable : {blog_path}"
        )

    workbook = load_workbook(
        blog_path,
        data_only=True,
        read_only=False
    )

    try:
        worksheet = workbook["RJ45"]

        # Détecte tous les blocs IN / OUT / STOCK
        components = find_component_blocks(
            worksheet
        )

        # Dernière ligne réelle du BLOG
        current_row = find_last_order_row(
            worksheet
        )

        results = []

        for component in components:

            stock = worksheet.cell(
                row=current_row,
                column=component["stock_col"]
            ).value

            try:
                stock = float(stock)
            except (TypeError, ValueError):
                stock = None

            name = str(
                component["name"]
            ).replace("\n", " ").strip()

            results.append({
                "component": name,
                "stock": stock,
                "stock_col": component["stock_col"]
            })

        return results
    finally:
        workbook.close()


def get_all_current_stock():
    if not BLOG_PATH.exists():
        raise FileNotFoundError(
            f"BLOG RJ45 introuvable : {BLOG_PATH}"
        )

    stat = BLOG_PATH.stat()
    cache_key = (
        str(BLOG_PATH.resolve()),
        stat.st_mtime_ns,
        stat.st_size
    )

    # Avoid loading the workbook twice if dashboard and stock open together.
    with _stock_cache_lock:
        results = _load_current_stock(*cache_key)

    return [dict(item) for item in results]