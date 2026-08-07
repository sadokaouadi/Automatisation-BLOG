from pathlib import Path
from openpyxl import load_workbook


# =========================================================
# OUTILS
# =========================================================

def normalize(value):
    if value is None:
        return ""

    return str(value).strip().replace(" ", "").upper()


def to_number(value):
    if value is None or value == "":
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# =========================================================
# DÉTECTION DES BLOCS COMPOSANTS
# =========================================================

def find_component_blocks(ws):
    """
    Détecte automatiquement les groupes :
    IN | OUT | STOCK
    """

    components = []

    # A:G = données principales.
    # Les composants commencent à partir de H.
    for stock_col in range(8, ws.max_column + 1):

        header = ws.cell(
            row=3,
            column=stock_col
        ).value

        if normalize(header) != "STOCK":
            continue

        in_col = stock_col - 2
        out_col = stock_col - 1

        component_name = ws.cell(
            row=1,
            column=in_col
        ).value

        if component_name in [None, ""]:
            continue

        components.append({
            "name": str(component_name).strip(),
            "in_col": in_col,
            "out_col": out_col,
            "stock_col": stock_col
        })

    return components


# =========================================================
# RECHERCHE DES LIGNES
# =========================================================

def find_latest_product_row(ws, product_code):
    """
    Dernière occurrence historique du produit.
    Produit = colonne D
    """

    target = normalize(product_code)

    for row in range(ws.max_row, 4, -1):

        value = ws.cell(
            row=row,
            column=4
        ).value

        if normalize(value) == target:
            return row

    return None


def find_last_order_row(ws):
    """
    Dernière ligne réelle contenant un produit.
    Cette ligne donne les stocks les plus récents.
    """

    for row in range(ws.max_row, 4, -1):

        product = ws.cell(
            row=row,
            column=4
        ).value

        if product not in [None, ""]:
            return row

    return None


# =========================================================
# ANALYSE DU STOCK
# =========================================================

def analyze_rj45_stock(
    blog_path,
    product_code,
    requested_qty
):

    blog_path = Path(blog_path)

    if not blog_path.exists():
        raise FileNotFoundError(
            f"BLOG introuvable : {blog_path}"
        )

    wb = load_workbook(
        blog_path,
        data_only=True
    )

    if "RJ45" not in wb.sheetnames:
        wb.close()
        raise ValueError(
            "La feuille RJ45 est introuvable."
        )

    ws = wb["RJ45"]

    # =====================================================
    # 1. Trouver l'historique du produit
    # =====================================================

    historical_row = find_latest_product_row(
        ws,
        product_code
    )

    if historical_row is None:

        wb.close()

        return {
            "found": False,
            "product": product_code,
            "status": "UNKNOWN",
            "message": "Produit introuvable dans l'historique.",
            "components": []
        }

    # =====================================================
    # 2. Quantité historique
    # =====================================================

    historical_qty = to_number(
        ws.cell(
            row=historical_row,
            column=5
        ).value
    )

    if historical_qty is None or historical_qty <= 0:

        wb.close()

        return {
            "found": False,
            "product": product_code,
            "status": "UNKNOWN",
            "message": "Quantité historique invalide.",
            "components": []
        }

    # =====================================================
    # 3. Ligne contenant le stock actuel
    # =====================================================

    current_stock_row = find_last_order_row(ws)

    if current_stock_row is None:

        wb.close()

        return {
            "found": False,
            "product": product_code,
            "status": "UNKNOWN",
            "message": "Impossible de trouver le stock actuel.",
            "components": []
        }

    # =====================================================
    # 4. Détecter les composants
    # =====================================================

    component_blocks = find_component_blocks(ws)

    components_used = []

    # =====================================================
    # 5. Analyser chaque composant utilisé
    # =====================================================

    for component in component_blocks:

        historical_out = to_number(
            ws.cell(
                row=historical_row,
                column=component["out_col"]
            ).value
        )

        # Pas consommé par ce produit
        if historical_out is None:
            continue

        if historical_out <= 0:
            continue

        # ---------------------------------------------
        # Consommation pour fabriquer 1 produit
        # ---------------------------------------------

        consumption_per_unit = (
            historical_out / historical_qty
        )

        # ---------------------------------------------
        # Stock actuel
        # ---------------------------------------------

        current_stock = to_number(
            ws.cell(
                row=current_stock_row,
                column=component["stock_col"]
            ).value
        )

        # ---------------------------------------------
        # Besoin pour nouvelle commande
        # ---------------------------------------------

        required_qty = (
            consumption_per_unit
            * float(requested_qty)
        )

        # ---------------------------------------------
        # Stock après commande
        # ---------------------------------------------

        if current_stock is None:

            projected_stock = None
            status = "UNKNOWN"

        else:

            projected_stock = (
                current_stock
                - required_qty
            )

            if projected_stock < 0:
                status = "INSUFFICIENT"
            else:
                status = "SUFFICIENT"

        components_used.append({
            "component": component["name"],

            "consumption_per_unit": round(
                consumption_per_unit,
                4
            ),

            "current_stock": (
                round(current_stock, 2)
                if current_stock is not None
                else None
            ),

            "required_qty": round(
                required_qty,
                2
            ),

            "projected_stock": (
                round(projected_stock, 2)
                if projected_stock is not None
                else None
            ),

            "status": status
        })

    wb.close()

    # =====================================================
    # 6. Statut global
    # =====================================================

    insufficient_components = [
        component
        for component in components_used
        if component["status"] == "INSUFFICIENT"
    ]

    unknown_components = [
        component
        for component in components_used
        if component["status"] == "UNKNOWN"
    ]

    if insufficient_components:

        global_status = "INSUFFICIENT"

    elif unknown_components:

        global_status = "UNKNOWN"

    else:

        global_status = "SUFFICIENT"

    # =====================================================
    # 7. Composant le plus critique
    # =====================================================

    valid_components = [
        component
        for component in components_used
        if component["projected_stock"] is not None
    ]

    limiting_component = None

    if valid_components:

        limiting_component = min(
            valid_components,
            key=lambda component:
                component["projected_stock"]
        )

    # =====================================================
    # 8. Résultat
    # =====================================================

    return {
        "found": True,

        "product": product_code,

        "requested_qty": float(
            requested_qty
        ),

        "historical_row": historical_row,

        "historical_qty": historical_qty,

        "current_stock_row":
            current_stock_row,

        "status":
            global_status,

        "components_count":
            len(components_used),

        "insufficient_count":
            len(insufficient_components),

        "limiting_component":
            limiting_component,

        "components":
            components_used
    }


# =========================================================
# TEST
# =========================================================

if __name__ == "__main__":

    result = analyze_rj45_stock(
        blog_path=(
            "data/input/"
            "BLOG 2026 - Rj45_cleaned.xlsx"
        ),
        product_code="NPC6ASZDB-WT003M",
        requested_qty=190000
    )

    print()
    print("==================================")
    print("ANALYSE STOCK RJ45")
    print("==================================")

    print(
        "Produit :",
        result.get("product")
    )

    print(
        "Quantité demandée :",
        result.get("requested_qty")
    )

    print(
        "Ligne historique :",
        result.get("historical_row")
    )

    print(
        "Quantité historique :",
        result.get("historical_qty")
    )

    print(
        "Ligne stock actuel :",
        result.get("current_stock_row")
    )

    print(
        "Statut global :",
        result.get("status")
    )

    print(
        "Nombre composants :",
        result.get("components_count")
    )

    print(
        "Composants insuffisants :",
        result.get("insufficient_count")
    )

    print()
    print("==================================")
    print("COMPOSANTS")
    print("==================================")

    for component in result.get(
        "components",
        []
    ):

        print()
        print("------------------------------")

        print(
            "Composant :",
            component["component"]
        )

        print(
            "Consommation / produit :",
            component["consumption_per_unit"]
        )

        print(
            "Stock actuel :",
            component["current_stock"]
        )

        print(
            "Besoin commande :",
            component["required_qty"]
        )

        print(
            "Stock après commande :",
            component["projected_stock"]
        )

        print(
            "Statut :",
            component["status"]
        )

    print()
    print("==================================")
    print("COMPOSANT LIMITANT")
    print("==================================")

    limiting = result.get(
        "limiting_component"
    )

    if limiting:

        print(
            "Composant :",
            limiting["component"]
        )

        print(
            "Stock actuel :",
            limiting["current_stock"]
        )

        print(
            "Besoin :",
            limiting["required_qty"]
        )

        print(
            "Stock après commande :",
            limiting["projected_stock"]
        )