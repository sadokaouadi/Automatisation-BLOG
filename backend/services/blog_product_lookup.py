from openpyxl import load_workbook
from pathlib import Path


def clean_value(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize(value):
    return clean_value(value).replace(" ", "").upper()


def find_product_in_rj45_blog(blog_path, product_code):
    """
    Cherche la dernière occurrence d'un produit dans le BLOG RJ45.

    Colonnes :
    D = produit
    F = mt
    G = UPC
    J = stock
    """

    blog_path = Path(blog_path)

    if not blog_path.exists():
        raise FileNotFoundError(f"Fichier introuvable : {blog_path}")

    wb = load_workbook(blog_path, data_only=True)
    ws = wb["RJ45"]

    product_target = normalize(product_code)
    results = []

    for row in range(5, ws.max_row + 1):
        product = normalize(ws.cell(row=row, column=4).value)

        if product == product_target:
            mt = ws.cell(row=row, column=6).value
            upc = ws.cell(row=row, column=7).value

            # Colonne J = 10
            stock = ws.cell(row=row, column=10).value

            results.append({
                "found": True,
                "sheet": "RJ45",
                "row": row,
                "product": product_code,
                "mt": mt,
                "upc": upc,
                "stock": stock
            })

    wb.close()

    if results:
        # Dernière occurrence du produit
        return results[-1]

    return {
        "found": False,
        "sheet": "RJ45",
        "row": None,
        "product": product_code,
        "mt": None,
        "upc": None,
        "stock": None
    }